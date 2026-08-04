# Copyright 2026 EPAM Systems, Inc. (“EPAM”)
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import csv
import io
import os
import re
from pathlib import Path
from typing import Any, List, Literal, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from codemie_tools.base.codemie_tool import CodeMieTool
from codemie_tools.data_management.code_executor.file_export_service import (
    FileExportService,
)
from codemie_tools.data_management.file_system.tools_vars import EXPORT_TABLES_TOOL

# Excel limits: sheet titles max 31 chars and may not contain these characters.
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_MAX_SHEET_TITLE = 31
_MAX_COLUMN_WIDTH = 60

# Excel structural limits: openpyxl silently writes a spec-invalid workbook past these.
_MAX_XLSX_ROWS = 1_048_575  # Excel hard limit minus header row
_MAX_XLSX_COLUMNS = 16_384  # Excel hard limit
_MAX_TOTAL_CELLS = 500_000  # DoS cap across all tables combined, both formats

# Control characters openpyxl rejects (matches openpyxl's own ILLEGAL_CHARACTERS_RE).
_ILLEGAL_CELL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
# Leading characters a spreadsheet interprets as a formula (CSV/formula injection).
_FORMULA_LEAD_CHARS = ("=", "+", "-", "@", "\t", "\r")


def _sanitize_cell(value: Any) -> Any:
    """Make a raw cell value safe to write to CSV or XLSX.

    Coerces non-primitive values to str (openpyxl rejects them), strips control
    characters openpyxl cannot serialize, and neutralizes leading characters that
    a spreadsheet would evaluate as a formula.
    """
    if value is None:
        return ""
    if not isinstance(value, (str, int, float, bool)):
        value = str(value)
    if isinstance(value, str):
        value = _ILLEGAL_CELL_CHARS.sub("", value)
        if value[:1] in _FORMULA_LEAD_CHARS:
            value = f"'{value}"
    return value


class TableSpec(BaseModel):
    sheet_name: str = Field(description="Sheet title for the .xlsx tab. Ignored for .csv content.")
    columns: List[str] = Field(description="Header row values, left to right.")
    rows: List[List[Any]] = Field(description="Data rows; each inner list is one row of cell values.")


class ExportTablesInput(BaseModel):
    tables: List[TableSpec] = Field(
        description="Tables to export. .xlsx writes one sheet per table; .csv needs exactly one."
    )
    filename: str = Field(description="Base file name for the download, e.g. 'sales_report'. Extension is optional.")
    format: Literal["xlsx", "csv"] = Field(description="Output format: 'xlsx' or 'csv'.")


class ExportTablesTool(CodeMieTool):
    name: str = EXPORT_TABLES_TOOL.name
    description: str = EXPORT_TABLES_TOOL.description or ""
    args_schema: Any = ExportTablesInput
    file_repository: Optional[Any] = Field(exclude=True, default=None)
    user_id: str = ""

    def execute(self, tables: Any, filename: str, format: str, *args, **kwargs) -> str:
        specs = [t if isinstance(t, TableSpec) else TableSpec(**t) for t in tables]
        if not specs:
            raise ValueError("No tables provided to export.")

        self._validate_size(specs, format)

        if format == "csv":
            if len(specs) > 1:
                raise ValueError("CSV format supports exactly one table; use 'xlsx' to export multiple tables.")
            content = self._build_csv(specs[0])
        elif format == "xlsx":
            content = self._build_xlsx(specs)
        else:
            raise ValueError(f"Unsupported export format: {format!r}")

        stored_filename = self._normalize_filename(filename, format)
        service = FileExportService(file_repository=self.file_repository, user_id=self.user_id)
        result = service.store_exported_bytes(stored_filename, content)
        if result is None:
            raise ValueError("File export is not configured.")
        return result

    @staticmethod
    def _validate_size(specs: List[TableSpec], fmt: str) -> None:
        total_cells = sum(len(t.rows) * max(len(t.columns), 1) for t in specs)
        if total_cells > _MAX_TOTAL_CELLS:
            raise ValueError(f"Export too large: {total_cells} cells requested, maximum is {_MAX_TOTAL_CELLS}.")

        if fmt != "xlsx":
            return

        for table in specs:
            if len(table.rows) > _MAX_XLSX_ROWS:
                raise ValueError(
                    f"Sheet {table.sheet_name!r} has {len(table.rows)} rows, "
                    f"exceeding the Excel limit of {_MAX_XLSX_ROWS} data rows."
                )
            if len(table.columns) > _MAX_XLSX_COLUMNS:
                raise ValueError(
                    f"Sheet {table.sheet_name!r} has {len(table.columns)} columns, "
                    f"exceeding the Excel limit of {_MAX_XLSX_COLUMNS} columns."
                )

    @staticmethod
    def _normalize_filename(filename: str, fmt: str) -> str:
        base = os.path.basename((filename or "").strip()) or "export"
        suffix = f".{fmt}"
        path = Path(base)
        if path.suffix.lower() == suffix:
            return base
        return f"{path.stem}{suffix}"

    @staticmethod
    def _build_csv(table: TableSpec) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow([_sanitize_cell(col) for col in table.columns])
        for row in table.rows:
            writer.writerow([_sanitize_cell(cell) for cell in row])
        return buffer.getvalue().encode("utf-8")

    def _build_xlsx(self, tables: List[TableSpec]) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.active)
        used_titles: set[str] = set()

        for table in tables:
            title = self._sheet_title(table.sheet_name, used_titles)
            worksheet = workbook.create_sheet(title=title)

            worksheet.append([_sanitize_cell(col) for col in table.columns])
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
            worksheet.freeze_panes = "A2"

            for row in table.rows:
                worksheet.append([_sanitize_cell(cell) for cell in row])

            self._apply_column_widths(worksheet, table)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _sheet_title(raw: str, used: set) -> str:
        cleaned = _INVALID_SHEET_CHARS.sub("", (raw or "").strip())
        cleaned = cleaned[:_MAX_SHEET_TITLE].strip() or "Sheet"

        candidate = cleaned
        suffix = 2
        while candidate in used:
            tail = f"_{suffix}"
            candidate = f"{cleaned[: _MAX_SHEET_TITLE - len(tail)]}{tail}"
            suffix += 1

        used.add(candidate)
        return candidate

    @staticmethod
    def _apply_column_widths(worksheet: Any, table: TableSpec) -> None:
        for index in range(len(table.columns)):
            longest = len(str(table.columns[index]))
            for row in table.rows:
                if index < len(row) and row[index] is not None:
                    longest = max(longest, len(str(row[index])))
            width = min(longest + 2, _MAX_COLUMN_WIDTH)
            worksheet.column_dimensions[get_column_letter(index + 1)].width = width
