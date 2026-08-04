# Copyright 2026 EPAM Systems, Inc. (“EPAM”)
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import csv
import io
from unittest import TestCase
from unittest.mock import MagicMock, patch

import openpyxl

from codemie_tools.data_management.file_system import export_tables_tool
from codemie_tools.data_management.file_system.export_tables_tool import (
    _MAX_TOTAL_CELLS,
    _MAX_XLSX_COLUMNS,
    ExportTablesTool,
    TableSpec,
)


def _stored():
    sf = MagicMock()
    sf.to_encoded_url.return_value = "encoded-id"
    return sf


def _tool():
    repo = MagicMock()
    repo.write_file.return_value = _stored()
    return ExportTablesTool(file_repository=repo, user_id="u1"), repo


class ExportCsvTests(TestCase):
    def test_csv_single_table_round_trips_with_escaping(self):
        tool, repo = _tool()
        table = TableSpec(
            sheet_name="Data",
            columns=["a", "b"],
            rows=[["x,y", 'she said "hi"'], ["line1\nline2", None]],
        )
        result = tool.execute(tables=[table], filename="out", format="csv")

        kwargs = repo.write_file.call_args.kwargs
        self.assertEqual(kwargs["mime_type"], "text/csv")
        self.assertTrue(kwargs["name"].endswith("out.csv"))
        self.assertEqual(kwargs["owner"], "u1")
        parsed = list(csv.reader(io.StringIO(kwargs["content"].decode("utf-8"))))
        self.assertEqual(parsed[0], ["a", "b"])
        self.assertEqual(parsed[1], ["x,y", 'she said "hi"'])
        self.assertEqual(parsed[2], ["line1\nline2", ""])
        self.assertIn("sandbox:/v1/files/encoded-id", result)

    def test_csv_rejects_multiple_tables(self):
        tool, _ = _tool()
        t = TableSpec(sheet_name="s", columns=["a"], rows=[["1"]])
        with self.assertRaises(ValueError):
            tool.execute(tables=[t, t], filename="out", format="csv")

    def test_empty_tables_raises(self):
        tool, _ = _tool()
        with self.assertRaises(ValueError):
            tool.execute(tables=[], filename="out", format="csv")

    def test_missing_repository_raises(self):
        tool = ExportTablesTool(file_repository=None, user_id="u1")
        t = TableSpec(sheet_name="s", columns=["a"], rows=[["1"]])
        with self.assertRaises(ValueError):
            tool.execute(tables=[t], filename="out", format="csv")


class ExportXlsxTests(TestCase):
    def _load(self, repo):
        content = repo.write_file.call_args.kwargs["content"]
        return openpyxl.load_workbook(io.BytesIO(content))

    def test_multi_sheet_headers_rows_and_styling(self):
        tool, repo = _tool()
        tables = [
            TableSpec(
                sheet_name="Sales",
                columns=["Region", "Total"],
                rows=[["EU", 10], ["US", 20]],
            ),
            TableSpec(sheet_name="Costs", columns=["Item", "Amount"], rows=[["Rent", 5]]),
        ]
        result = tool.execute(tables=tables, filename="report", format="xlsx")

        kwargs = repo.write_file.call_args.kwargs
        self.assertEqual(
            kwargs["mime_type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertTrue(kwargs["name"].endswith("report.xlsx"))
        self.assertIn("sandbox:/v1/files/encoded-id", result)

        wb = self._load(repo)
        self.assertEqual(wb.sheetnames, ["Sales", "Costs"])
        ws = wb["Sales"]
        self.assertEqual([c.value for c in ws[1]], ["Region", "Total"])
        self.assertEqual(ws["A2"].value, "EU")
        self.assertEqual(ws["B3"].value, 20)
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws.freeze_panes, "A2")

    def test_sheet_name_sanitized_and_deduped(self):
        tool, repo = _tool()
        bad = "a" * 40 + "[x]:*?/\\"
        tables = [
            TableSpec(sheet_name=bad, columns=["c"], rows=[["1"]]),
            TableSpec(sheet_name=bad, columns=["c"], rows=[["2"]]),
        ]
        tool.execute(tables=tables, filename="r", format="xlsx")

        wb = self._load(repo)
        self.assertEqual(len(wb.sheetnames), 2)
        self.assertEqual(len(set(wb.sheetnames)), 2)
        for name in wb.sheetnames:
            self.assertLessEqual(len(name), 31)
            self.assertFalse(set(name) & set('[]:*?/\\'))

    def test_blank_sheet_name_falls_back(self):
        tool, repo = _tool()
        tables = [TableSpec(sheet_name="   ", columns=["c"], rows=[["1"]])]
        tool.execute(tables=tables, filename="r", format="xlsx")

        wb = self._load(repo)
        self.assertEqual(len(wb.sheetnames), 1)
        self.assertTrue(wb.sheetnames[0])

    def test_header_only_table_is_valid(self):
        tool, repo = _tool()
        tables = [TableSpec(sheet_name="Empty", columns=["a", "b"], rows=[])]
        tool.execute(tables=tables, filename="r", format="xlsx")

        wb = self._load(repo)
        self.assertEqual([c.value for c in wb["Empty"][1]], ["a", "b"])
        self.assertEqual(wb["Empty"].max_row, 1)

    def test_large_table_is_valid(self):
        tool, repo = _tool()
        rows = [[i, f"row{i}"] for i in range(3000)]
        tables = [TableSpec(sheet_name="Big", columns=["n", "label"], rows=rows)]
        tool.execute(tables=tables, filename="r", format="xlsx")

        wb = self._load(repo)
        self.assertEqual(wb["Big"].max_row, 3001)

    def test_non_primitive_cells_are_stringified(self):
        # openpyxl raises on dict/list cells; the tool must coerce them (like csv does).
        tool, repo = _tool()
        tables = [
            TableSpec(
                sheet_name="S",
                columns=["obj", "arr"],
                rows=[[{"k": 1}, [1, 2]]],
            )
        ]
        tool.execute(tables=tables, filename="r", format="xlsx")
        wb = self._load(repo)
        ws = wb["S"]
        self.assertEqual(ws["A2"].value, str({"k": 1}))
        self.assertEqual(ws["B2"].value, str([1, 2]))

    def test_illegal_control_chars_stripped(self):
        # openpyxl raises IllegalCharacterError on control chars; they must be removed.
        tool, repo = _tool()
        tables = [TableSpec(sheet_name="S", columns=["c"], rows=[["a\x00b\x07c"]])]
        tool.execute(tables=tables, filename="r", format="xlsx")
        wb = self._load(repo)
        self.assertEqual(wb["S"]["A2"].value, "abc")

    def test_native_numeric_cells_preserved(self):
        tool, repo = _tool()
        tables = [TableSpec(sheet_name="S", columns=["n", "f", "b"], rows=[[7, 1.5, True]])]
        tool.execute(tables=tables, filename="r", format="xlsx")
        ws = self._load(repo)["S"]
        self.assertEqual(ws["A2"].value, 7)
        self.assertEqual(ws["B2"].value, 1.5)
        self.assertIs(ws["C2"].value, True)

    def test_xlsx_neutralizes_formula_injection(self):
        tool, repo = _tool()
        tables = [TableSpec(sheet_name="S", columns=["x"], rows=[["=1+2"], ["@cmd"], ["-9"]])]
        tool.execute(tables=tables, filename="r", format="xlsx")
        ws = self._load(repo)["S"]
        self.assertEqual(ws["A2"].value, "'=1+2")
        self.assertEqual(ws["A3"].value, "'@cmd")
        self.assertEqual(ws["A4"].value, "'-9")


class ExportSizeLimitTests(TestCase):
    def test_xlsx_rejects_too_many_rows(self):
        # The real row cap (1_048_575) can only be reached via inputs that already
        # blow the total-cell cap, so the guard is exercised with a lowered limit.
        tool, repo = _tool()
        rows = [[i] for i in range(5)]
        tables = [TableSpec(sheet_name="TooTall", columns=["n"], rows=rows)]
        with patch.object(export_tables_tool, "_MAX_XLSX_ROWS", 4):
            with self.assertRaises(ValueError) as ctx:
                tool.execute(tables=tables, filename="r", format="xlsx")
        self.assertIn("TooTall", str(ctx.exception))
        repo.write_file.assert_not_called()

    def test_xlsx_rejects_too_many_columns(self):
        tool, repo = _tool()
        columns = [f"c{i}" for i in range(_MAX_XLSX_COLUMNS + 1)]
        tables = [TableSpec(sheet_name="TooWide", columns=columns, rows=[])]
        with self.assertRaises(ValueError) as ctx:
            tool.execute(tables=tables, filename="r", format="xlsx")
        self.assertIn("TooWide", str(ctx.exception))
        repo.write_file.assert_not_called()

    def test_xlsx_rejects_total_cells_over_cap(self):
        tool, repo = _tool()
        columns = ["a", "b", "c", "d"]
        rows = [[1, 2, 3, 4]] * ((_MAX_TOTAL_CELLS // 4) + 1)
        tables = [TableSpec(sheet_name="S", columns=columns, rows=rows)]
        with self.assertRaises(ValueError) as ctx:
            tool.execute(tables=tables, filename="r", format="xlsx")
        self.assertIn(str(_MAX_TOTAL_CELLS), str(ctx.exception))
        repo.write_file.assert_not_called()

    def test_csv_rejects_total_cells_over_cap(self):
        tool, repo = _tool()
        columns = ["a", "b", "c", "d"]
        rows = [[1, 2, 3, 4]] * ((_MAX_TOTAL_CELLS // 4) + 1)
        tables = [TableSpec(sheet_name="S", columns=columns, rows=rows)]
        with self.assertRaises(ValueError) as ctx:
            tool.execute(tables=tables, filename="r", format="csv")
        self.assertIn(str(_MAX_TOTAL_CELLS), str(ctx.exception))
        repo.write_file.assert_not_called()

    def test_total_cells_summed_across_tables(self):
        tool, repo = _tool()
        half = [[1, 2]] * (_MAX_TOTAL_CELLS // 4)
        tables = [
            TableSpec(sheet_name="A", columns=["x", "y"], rows=half),
            TableSpec(sheet_name="B", columns=["x", "y"], rows=half + [[1, 2]]),
        ]
        with self.assertRaises(ValueError):
            tool.execute(tables=tables, filename="r", format="xlsx")
        repo.write_file.assert_not_called()

    def test_just_under_total_cell_cap_succeeds(self):
        tool, repo = _tool()
        rows = [[i, i] for i in range(_MAX_TOTAL_CELLS // 2)]
        tables = [TableSpec(sheet_name="S", columns=["a", "b"], rows=rows)]
        tool.execute(tables=tables, filename="r", format="xlsx")

        content = repo.write_file.call_args.kwargs["content"]
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        self.assertEqual(wb["S"].max_row, (_MAX_TOTAL_CELLS // 2) + 1)
        wb.close()


class SanitizeCellTests(TestCase):
    def test_csv_neutralizes_formula_injection(self):
        tool, repo = _tool()
        table = TableSpec(sheet_name="s", columns=["x"], rows=[["=SUM(A1:A2)"], ["+1"]])
        tool.execute(tables=[table], filename="out", format="csv")
        content = repo.write_file.call_args.kwargs["content"].decode("utf-8")
        parsed = list(csv.reader(io.StringIO(content)))
        self.assertEqual(parsed[1], ["'=SUM(A1:A2)"])
        self.assertEqual(parsed[2], ["'+1"])

    def test_csv_leaves_safe_and_numeric_values(self):
        tool, repo = _tool()
        table = TableSpec(sheet_name="s", columns=["x", "n"], rows=[["hello", 5]])
        tool.execute(tables=[table], filename="out", format="csv")
        content = repo.write_file.call_args.kwargs["content"].decode("utf-8")
        parsed = list(csv.reader(io.StringIO(content)))
        self.assertEqual(parsed[1], ["hello", "5"])
